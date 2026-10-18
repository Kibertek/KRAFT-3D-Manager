import hashlib
import sqlite3
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QTableWidget, QTableWidgetItem, QHeaderView, QMessageBox,
    QFileDialog, QAbstractItemView
)

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "data" / "kraft3d.db"


def get_connection():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def money(value):
    return f"{float(value or 0):,.2f} ₽".replace(",", " ").replace(".", ",")


def number(value):
    if value is None:
        return 0.0
    text = str(value).strip().replace("\xa0", "").replace(" ", "")
    if not text:
        return 0.0
    text = text.replace("%", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def file_hash(path):
    digest = hashlib.sha256()
    with open(path, "rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def init_ozon_finance_tables():
    with get_connection() as connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS marketplace_accruals (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                marketplace TEXT NOT NULL,
                import_id INTEGER,
                accrual_id TEXT NOT NULL,
                accrual_date TEXT,
                service_group TEXT,
                accrual_type TEXT,
                article TEXT,
                sku TEXT,
                product_name TEXT,
                quantity REAL DEFAULT 0,
                seller_price REAL DEFAULT 0,
                service_date TEXT,
                sales_platform TEXT,
                work_scheme TEXT,
                commission_percent REAL DEFAULT 0,
                localization_index_percent REAL DEFAULT 0,
                delivery_time_hours REAL DEFAULT 0,
                amount_total REAL DEFAULT 0,
                product_id INTEGER,
                raw_data TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(marketplace, import_id, accrual_id, id),
                FOREIGN KEY(import_id) REFERENCES marketplace_imports(id) ON DELETE SET NULL,
                FOREIGN KEY(product_id) REFERENCES products(id) ON DELETE SET NULL
            )
        """)
        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_marketplace_accruals_article
            ON marketplace_accruals(article)
        """)
        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_marketplace_accruals_date
            ON marketplace_accruals(accrual_date)
        """)
        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_marketplace_accruals_accrual
            ON marketplace_accruals(accrual_id)
        """)


def find_header(rows):
    required = {
        "ID начисления",
        "Группа услуг",
        "Тип начисления",
        "Сумма итого, руб.",
    }
    for index, row in enumerate(rows):
        headers = {clean(value) for value in row}
        if required.issubset(headers):
            return index
    return None


def parse_ozon_financial_report(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    header_index = find_header(rows)
    if header_index is None:
        raise ValueError(
            "Не удалось найти таблицу финансового отчёта Ozon. "
            "Проверь, что выбран файл «Отчёт о начислениях»."
        )

    header = [clean(value) for value in rows[header_index]]
    positions = {name: index for index, name in enumerate(header) if name}

    def get(row, *names):
        for name in names:
            index = positions.get(name)
            if index is not None and index < len(row):
                return row[index]
        return None

    data = []
    for row in rows[header_index + 1:]:
        accrual_id = clean(get(row, "ID начисления"))
        service_group = clean(get(row, "Группа услуг"))
        accrual_type = clean(get(row, "Тип начисления"))
        amount = get(row, "Сумма итого, руб.")

        if not accrual_id and not service_group and not accrual_type:
            continue

        if not accrual_id:
            continue

        data.append({
            "accrual_id": accrual_id,
            "accrual_date": clean(get(row, "Дата начисления")),
            "service_group": service_group,
            "accrual_type": accrual_type,
            "article": clean(get(row, "Артикул")),
            "sku": clean(get(row, "SKU")),
            "product_name": clean(get(row, "Название товара")),
            "quantity": number(get(row, "Количество")),
            "seller_price": number(get(row, "Цена продавца")),
            "service_date": clean(get(
                row,
                "Дата принятия заказа в обработку или оказания услуги",
                "Дата принятия заказа в обработку или оказания услуги ",
            )),
            "sales_platform": clean(get(row, "Платформа продажи")),
            "work_scheme": clean(get(row, "Схема работы")),
            "commission_percent": number(get(row, "Вознаграждение Ozon, %")),
            "localization_index_percent": number(get(row, "Индекс локализации, %")),
            "delivery_time_hours": number(get(row, "Среднее время доставки, часы")),
            "amount_total": number(amount),
        })

    return data


def import_ozon_financial_report(path):
    rows = parse_ozon_financial_report(path)
    digest = file_hash(path)

    with get_connection() as connection:
        existing = connection.execute(
            "SELECT id FROM marketplace_imports WHERE file_hash = ?",
            (digest,)
        ).fetchone()
        if existing:
            return 0, 0, "Этот файл уже импортирован."

        cursor = connection.execute("""
            INSERT INTO marketplace_imports (
                marketplace, report_type, file_name, file_hash, rows_count
            ) VALUES (?, ?, ?, ?, 0)
        """, (
            "Ozon",
            "Финансовый отчёт",
            path.name,
            digest,
        ))
        import_id = cursor.lastrowid

        added = 0
        for row in rows:
            product = None
            if row["article"]:
                product = connection.execute(
                    "SELECT id FROM products WHERE article = ? COLLATE NOCASE",
                    (row["article"],)
                ).fetchone()
            product_id = product["id"] if product else None

            connection.execute("""
                INSERT INTO marketplace_accruals (
                    marketplace, import_id, accrual_id, accrual_date,
                    service_group, accrual_type, article, sku, product_name,
                    quantity, seller_price, service_date, sales_platform,
                    work_scheme, commission_percent,
                    localization_index_percent, delivery_time_hours,
                    amount_total, product_id, raw_data
                ) VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                )
            """, (
                "Ozon", import_id, row["accrual_id"], row["accrual_date"],
                row["service_group"], row["accrual_type"], row["article"],
                row["sku"], row["product_name"], row["quantity"],
                row["seller_price"], row["service_date"], row["sales_platform"],
                row["work_scheme"], row["commission_percent"],
                row["localization_index_percent"], row["delivery_time_hours"],
                row["amount_total"], product_id, repr(row),
            ))
            added += 1

        connection.execute(
            "UPDATE marketplace_imports SET rows_count = ? WHERE id = ?",
            (added, import_id),
        )

    return added, 0, None


def is_revenue_row(service_group, accrual_type, amount):
    if amount <= 0:
        return False
    if service_group == "Продажи":
        return True
    return accrual_type in {
        "Выручка",
        "Программы партнёров",
        "Баллы за скидки",
    }


def is_global_expense(service_group, accrual_type, article):
    if article:
        return False
    return service_group in {
        "Продвижение и реклама",
        "Другие услуги и штрафы",
    } or accrual_type in {
        "Досрочная выплата",
    }


def load_profit_rows():
    with get_connection() as connection:
        rows = connection.execute("""
            SELECT
                article,
                MAX(product_name) AS product_name,
                MAX(product_id) AS product_id,
                SUM(CASE
                    WHEN amount_total > 0 AND (
                        service_group = 'Продажи'
                        OR accrual_type IN (
                            'Выручка',
                            'Программы партнёров',
                            'Баллы за скидки'
                        )
                    ) THEN amount_total ELSE 0 END) AS gross_revenue,
                SUM(CASE WHEN amount_total < 0 THEN -amount_total ELSE 0 END)
                    AS marketplace_expenses,
                SUM(CASE WHEN amount_total < 0 AND service_group = 'Вознаграждение Ozon'
                    THEN -amount_total ELSE 0 END) AS ozon_commission,
                SUM(CASE WHEN amount_total < 0 AND service_group = 'Услуги доставки'
                    THEN -amount_total ELSE 0 END) AS delivery_expenses,
                SUM(CASE WHEN amount_total < 0 AND service_group = 'Услуги партнёров'
                    THEN -amount_total ELSE 0 END) AS partner_expenses,
                SUM(CASE WHEN amount_total < 0 AND service_group = 'Продвижение и реклама'
                    THEN -amount_total ELSE 0 END) AS product_ad_expenses
            FROM marketplace_accruals
            WHERE marketplace = 'Ozon'
              AND article <> ''
            GROUP BY article
            ORDER BY gross_revenue DESC, article
        """).fetchall()

        global_expenses = connection.execute("""
            SELECT COALESCE(SUM(-amount_total), 0) AS total
            FROM marketplace_accruals
            WHERE marketplace = 'Ozon'
              AND amount_total < 0
              AND article = ''
        """).fetchone()["total"]

        global_by_type = connection.execute("""
            SELECT service_group, accrual_type,
                   COALESCE(SUM(-amount_total), 0) AS total
            FROM marketplace_accruals
            WHERE marketplace = 'Ozon'
              AND amount_total < 0
              AND article = ''
            GROUP BY service_group, accrual_type
            ORDER BY total DESC
        """).fetchall()

        return rows, global_expenses, global_by_type


class OzonFinancePage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        init_ozon_finance_tables()
        self.setup_ui()
        self.load_data()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(35, 30, 35, 30)
        layout.setSpacing(14)

        title = QLabel("Финансы Ozon")
        title.setObjectName("pageTitle")
        layout.addWidget(title)

        subtitle = QLabel(
            "Финансовый отчёт Ozon разбирается по начислениям: выручка, "
            "комиссия, логистика, услуги партнёров и реклама. Себестоимость "
            "берётся из каталога KRAFT 3D, налог пока считается отдельно."
        )
        subtitle.setObjectName("subtitle")
        subtitle.setWordWrap(True)
        layout.addWidget(subtitle)

        buttons = QHBoxLayout()
        import_button = QPushButton("Импорт финансового отчёта Ozon")
        import_button.setObjectName("accentButton")
        import_button.clicked.connect(self.import_report)
        buttons.addWidget(import_button)

        refresh_button = QPushButton("Обновить")
        refresh_button.clicked.connect(self.load_data)
        buttons.addWidget(refresh_button)
        buttons.addStretch()
        layout.addLayout(buttons)

        self.summary = QLabel("")
        self.summary.setObjectName("pageTitle")
        self.summary.setWordWrap(True)
        layout.addWidget(self.summary)

        self.table = QTableWidget(0, 9)
        self.table.setHorizontalHeaderLabels([
            "Артикул", "Товар", "Выручка", "Комиссия Ozon",
            "Логистика", "Партнёры", "Реклама", "Себестоимость", "Прибыль"
        ])
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        header.setStretchLastSection(True)
        layout.addWidget(self.table, 1)

    def import_report(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Выбери финансовый отчёт Ozon",
            str(Path.home()),
            "Excel (*.xlsx)",
        )
        if not path:
            return

        try:
            added, duplicates, message = import_ozon_financial_report(Path(path))
            if message:
                QMessageBox.information(self, "Импорт Ozon", message)
                return
            QMessageBox.information(
                self,
                "Импорт завершён",
                f"Добавлено начислений: {added}\nПропущено дублей: {duplicates}",
            )
            self.load_data()
        except Exception as error:
            QMessageBox.critical(
                self,
                "Ошибка импорта Ozon",
                f"Не удалось импортировать финансовый отчёт:\n{error}",
            )

    def load_data(self):
        rows, global_expenses, global_by_type = load_profit_rows()

        self.table.setRowCount(len(rows))
        total_revenue = 0.0
        total_expenses = 0.0
        total_cost = 0.0
        total_profit = 0.0

        with get_connection() as connection:
            for row_index, row in enumerate(rows):
                product = None
                if row["product_id"]:
                    product = connection.execute(
                        "SELECT material_cost, components_cost, production_cost "
                        "FROM products WHERE id = ?",
                        (row["product_id"],),
                    ).fetchone()

                quantity = connection.execute("""
                    SELECT COALESCE(SUM(quantity), 0) AS quantity
                    FROM marketplace_accruals
                    WHERE marketplace = 'Ozon'
                      AND article = ?
                      AND service_group = 'Продажи'
                      AND quantity > 0
                """, (row["article"],)).fetchone()["quantity"]

                direct_unit_cost = 0.0
                if product:
                    direct_unit_cost = (
                        float(product["material_cost"] or 0)
                        + float(product["components_cost"] or 0)
                        + float(product["production_cost"] or 0)
                    )

                product_cost = direct_unit_cost * float(quantity or 0)
                expenses = float(row["marketplace_expenses"] or 0)
                profit_before_tax = float(row["gross_revenue"] or 0) - expenses - product_cost

                values = [
                    row["article"] or "—",
                    row["product_name"] or "—",
                    money(row["gross_revenue"]),
                    money(row["ozon_commission"]),
                    money(row["delivery_expenses"]),
                    money(row["partner_expenses"]),
                    money(row["product_ad_expenses"]),
                    money(product_cost),
                    money(profit_before_tax),
                ]

                for column, value in enumerate(values):
                    item = QTableWidgetItem(str(value))
                    if column >= 2:
                        item.setTextAlignment(
                            Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter
                        )
                    self.table.setItem(row_index, column, item)

                total_revenue += float(row["gross_revenue"] or 0)
                total_expenses += expenses
                total_cost += product_cost
                total_profit += profit_before_tax

        global_text = " | ".join(
            f"{row['service_group']} / {row['accrual_type']}: {money(row['total'])}"
            for row in global_by_type
        )
        if not global_text:
            global_text = "нет"

        self.summary.setText(
            f"Продажи: {money(total_revenue)} | "
            f"Расходы Ozon по товарам: {money(total_expenses)} | "
            f"Себестоимость: {money(total_cost)} | "
            f"Прибыль до налога: {money(total_profit)}\n"
            f"Глобальные расходы без артикула: {money(global_expenses)}\n"
            f"{global_text}"
        )
