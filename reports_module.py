import sqlite3
from pathlib import Path

from PySide6.QtCore import Qt, QDate
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QDateEdit,
    QTableWidget, QTableWidgetItem, QHeaderView, QMessageBox, QFileDialog
)

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "data" / "kraft3d.db"


def get_connection():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def money(value):
    return f"{float(value or 0):,.2f} ₽".replace(",", " ").replace(".", ",")


class ReportsPage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setup_ui()
        self.load_report()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(35, 30, 35, 30)
        layout.setSpacing(14)

        title = QLabel("Отчёты")
        title.setObjectName("pageTitle")
        layout.addWidget(title)

        subtitle = QLabel(
            "Финансовый результат считается по фактически внесённым операциям. "
            "Суммы заказов показываются отдельно и не считаются полученным доходом автоматически."
        )
        subtitle.setObjectName("subtitle")
        subtitle.setWordWrap(True)
        layout.addWidget(subtitle)

        filters = QHBoxLayout()
        filters.addWidget(QLabel("С даты:"))
        self.start_date = QDateEdit()
        self.start_date.setCalendarPopup(True)
        self.start_date.setDisplayFormat("dd.MM.yyyy")
        self.start_date.setDate(QDate.currentDate().addMonths(-1).addDays(1))
        filters.addWidget(self.start_date)

        filters.addWidget(QLabel("По дату:"))
        self.end_date = QDateEdit()
        self.end_date.setCalendarPopup(True)
        self.end_date.setDisplayFormat("dd.MM.yyyy")
        self.end_date.setDate(QDate.currentDate())
        filters.addWidget(self.end_date)

        refresh_button = QPushButton("Сформировать")
        refresh_button.setObjectName("accentButton")
        refresh_button.clicked.connect(self.load_report)
        filters.addWidget(refresh_button)

        export_button = QPushButton("Экспорт в Excel")
        export_button.clicked.connect(self.export_excel)
        filters.addWidget(export_button)
        filters.addStretch()
        layout.addLayout(filters)

        self.summary_label = QLabel("Доходы: 0 ₽ | Расходы: 0 ₽ | Результат: 0 ₽")
        self.summary_label.setObjectName("pageTitle")
        self.summary_label.setWordWrap(True)
        layout.addWidget(self.summary_label)

        self.orders_summary_label = QLabel("Заказы за период: 0 | Сумма заказов: 0 ₽")
        self.orders_summary_label.setObjectName("subtitle")
        layout.addWidget(self.orders_summary_label)

        self.category_table = QTableWidget(0, 3)
        self.category_table.setHorizontalHeaderLabels(["Тип", "Категория", "Сумма"])
        self.category_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.category_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.category_table.verticalHeader().setVisible(False)
        self.category_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.ResizeToContents
        )
        self.category_table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(QLabel("Финансы по категориям"))
        layout.addWidget(self.category_table, 1)

        self.status_table = QTableWidget(0, 3)
        self.status_table.setHorizontalHeaderLabels(["Статус заказа", "Количество", "Сумма заказов"])
        self.status_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.status_table.verticalHeader().setVisible(False)
        self.status_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.ResizeToContents
        )
        self.status_table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(QLabel("Заказы по статусам"))
        layout.addWidget(self.status_table, 1)

        self.finance_rows = []
        self.order_rows = []
        self.category_rows = []
        self.status_rows = []

    def period(self):
        start = self.start_date.date().toString("yyyy-MM-dd")
        end = self.end_date.date().toString("yyyy-MM-dd")
        if start > end:
            raise ValueError("Начало периода не может быть позже его окончания.")
        return start, end

    def load_report(self, *_):
        try:
            start, end = self.period()
        except ValueError as error:
            QMessageBox.warning(self, "Неверный период", str(error))
            return

        with get_connection() as connection:
            self.finance_rows = [dict(row) for row in connection.execute("""
                SELECT operation_date, operation_type, category, amount, source,
                       order_number, marketplace, comment
                FROM finance_transactions
                WHERE operation_date BETWEEN ? AND ?
                ORDER BY operation_date, id
            """, (start, end)).fetchall()]

            self.order_rows = [dict(row) for row in connection.execute("""
                SELECT order_number, order_date, customer_name, product_name,
                       quantity, unit_price, total_price, status, due_date, comment
                FROM orders
                WHERE order_date BETWEEN ? AND ?
                ORDER BY order_date, id
            """, (start, end)).fetchall()]

        self.category_rows = []
        category_map = {}
        income = 0.0
        expense = 0.0
        for row in self.finance_rows:
            amount = float(row.get("amount") or 0)
            kind = row.get("operation_type") or "Расход"
            if kind == "Доход":
                income += amount
            else:
                expense += amount
            key = (kind, row.get("category") or "Без категории")
            category_map[key] = category_map.get(key, 0.0) + amount

        self.category_rows = [
            {"type": key[0], "category": key[1], "amount": amount}
            for key, amount in sorted(category_map.items(), key=lambda item: (item[0][0], item[0][1].lower()))
        ]

        self.status_rows = []
        status_map = {}
        for row in self.order_rows:
            key = row.get("status") or "Без статуса"
            values = status_map.setdefault(key, {"count": 0, "total": 0.0})
            values["count"] += 1
            values["total"] += float(row.get("total_price") or 0)
        self.status_rows = [
            {"status": key, **value}
            for key, value in sorted(status_map.items(), key=lambda item: item[0].lower())
        ]

        self.summary_label.setText(
            f"Фактические доходы: {money(income)}    |    "
            f"Фактические расходы: {money(expense)}    |    "
            f"Разница: {money(income - expense)}"
        )
        order_total = sum(float(row.get("total_price") or 0) for row in self.order_rows)
        self.orders_summary_label.setText(
            f"Заказы за период: {len(self.order_rows)}    |    "
            f"Сумма заказов: {money(order_total)}"
        )

        self.category_table.setRowCount(len(self.category_rows))
        for index, row in enumerate(self.category_rows):
            for col, value in enumerate((row["type"], row["category"], money(row["amount"]))):
                self.category_table.setItem(index, col, QTableWidgetItem(str(value)))

        self.status_table.setRowCount(len(self.status_rows))
        for index, row in enumerate(self.status_rows):
            values = (row["status"], str(row["count"]), money(row["total"]))
            for col, value in enumerate(values):
                item = QTableWidgetItem(value)
                if col > 0:
                    item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.status_table.setItem(index, col, item)

    def export_excel(self):
        try:
            start, end = self.period()
        except ValueError as error:
            QMessageBox.warning(self, "Неверный период", str(error))
            return

        default_name = f"KRAFT3D_Отчёт_{start}_{end}.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить отчёт Excel", str(BASE_DIR / default_name),
            "Книга Excel (*.xlsx)"
        )
        if not file_path:
            return
        if not file_path.lower().endswith(".xlsx"):
            file_path += ".xlsx"

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            summary = workbook.active
            summary.title = "Сводка"
            summary.append(["KRAFT 3D Manager", "Отчёт за период"])
            summary.append(["Начало периода", start])
            summary.append(["Конец периода", end])

            income = sum(float(row.get("amount") or 0) for row in self.finance_rows if row.get("operation_type") == "Доход")
            expense = sum(float(row.get("amount") or 0) for row in self.finance_rows if row.get("operation_type") != "Доход")
            order_total = sum(float(row.get("total_price") or 0) for row in self.order_rows)
            summary.append(["Фактические доходы", income])
            summary.append(["Фактические расходы", expense])
            summary.append(["Разница доходов и расходов", income - expense])
            summary.append(["Количество финансовых операций", len(self.finance_rows)])
            summary.append(["Количество заказов", len(self.order_rows)])
            summary.append(["Сумма заказов (не равна автоматически полученному доходу)", order_total])

            categories = workbook.create_sheet("По категориям")
            categories.append(["Тип", "Категория", "Сумма, ₽"])
            for row in self.category_rows:
                categories.append([row["type"], row["category"], row["amount"]])

            finance = workbook.create_sheet("Финансовые операции")
            finance.append(["Дата", "Тип", "Категория", "Сумма, ₽", "Источник", "Номер заказа", "Маркетплейс", "Комментарий"])
            for row in self.finance_rows:
                finance.append([
                    row.get("operation_date", ""), row.get("operation_type", ""),
                    row.get("category", ""), float(row.get("amount") or 0),
                    row.get("source", ""), row.get("order_number", ""),
                    row.get("marketplace", ""), row.get("comment", "")
                ])

            orders = workbook.create_sheet("Заказы")
            orders.append(["Номер", "Дата", "Заказчик", "Изделие", "Количество", "Цена за единицу, ₽", "Сумма, ₽", "Статус", "Срок выполнения", "Комментарий"])
            for row in self.order_rows:
                orders.append([
                    row.get("order_number", ""), row.get("order_date", ""),
                    row.get("customer_name", ""), row.get("product_name", ""),
                    row.get("quantity", 0), row.get("unit_price", 0),
                    row.get("total_price", 0), row.get("status", ""),
                    row.get("due_date", ""), row.get("comment", "")
                ])

            for sheet in workbook.worksheets:
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="303030")
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                for column in sheet.columns:
                    max_length = 0
                    column_letter = get_column_letter(column[0].column)
                    for cell in column:
                        value = "" if cell.value is None else str(cell.value)
                        max_length = max(max_length, len(value))
                        if cell.row > 1 and isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##0.00'
                    sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)

            workbook.save(file_path)
        except ImportError:
            QMessageBox.critical(
                self, "Не найден openpyxl",
                "Для экспорта в Excel установи библиотеку командой:\npy -m pip install openpyxl"
            )
            return
        except Exception as error:
            QMessageBox.critical(self, "Ошибка экспорта", f"Не удалось сохранить отчёт:\n{error}")
            return

        QMessageBox.information(self, "Отчёт сохранён", f"Excel-отчёт сохранён:\n{file_path}")
