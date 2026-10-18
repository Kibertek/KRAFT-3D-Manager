import csv
import hashlib
import sqlite3
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QTableWidget, QTableWidgetItem, QHeaderView, QMessageBox,
    QFileDialog, QAbstractItemView
)

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "data" / "kraft3d.db"


def get_connection():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def money(value):
    return f"{float(value or 0):,.2f} ₽".replace(",", " ").replace(".", ",")


def init_marketplace_tables():
    with get_connection() as connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS marketplace_imports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                marketplace TEXT NOT NULL,
                report_type TEXT NOT NULL,
                file_name TEXT NOT NULL,
                file_hash TEXT NOT NULL UNIQUE,
                imported_at TEXT DEFAULT CURRENT_TIMESTAMP,
                rows_count INTEGER DEFAULT 0
            )
        """)
        connection.execute("""
            CREATE TABLE IF NOT EXISTS marketplace_sales (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                marketplace TEXT NOT NULL,
                report_type TEXT NOT NULL,
                import_id INTEGER,
                external_key TEXT NOT NULL,
                order_number TEXT,
                shipment_number TEXT,
                sale_date TEXT,
                status TEXT,
                product_name TEXT,
                article TEXT,
                sku TEXT,
                barcode TEXT,
                quantity REAL DEFAULT 0,
                sale_price REAL DEFAULT 0,
                revenue REAL DEFAULT 0,
                loyalty_payout REAL DEFAULT 0,
                return_amount REAL DEFAULT 0,
                buyer_paid REAL DEFAULT 0,
                discount_percent REAL DEFAULT 0,
                discount_amount REAL DEFAULT 0,
                product_id INTEGER,
                raw_data TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(marketplace, report_type, external_key),
                FOREIGN KEY(import_id) REFERENCES marketplace_imports(id) ON DELETE SET NULL,
                FOREIGN KEY(product_id) REFERENCES products(id) ON DELETE SET NULL
            )
        """)
        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_marketplace_sales_article
            ON marketplace_sales(article)
        """)
        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_marketplace_sales_date
            ON marketplace_sales(sale_date)
        """)


def _file_hash(path):
    digest = hashlib.sha256()
    with open(path, "rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _number(value):
    if value is None:
        return 0.0
    text = str(value).strip().replace("\xa0", "").replace(" ", "")
    if not text:
        return 0.0
    text = text.replace("%", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def _clean(value):
    if value is None:
        return ""
    return str(value).strip()


def _find_products(connection, article):
    if not article:
        return None
    return connection.execute(
        "SELECT id FROM products WHERE article = ? COLLATE NOCASE",
        (article,)
    ).fetchone()


def _parse_ozon_realization(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    header_row = None
    for index, row in enumerate(rows):
        values = [_clean(value) for value in row]
        if "Название товара" in values and "Артикул" in values:
            header_row = index
            break

    if header_row is None:
        raise ValueError("Не удалось найти таблицу отчёта о реализации Ozon.")

    data = []
    for row in rows[header_row + 3:]:
        values = list(row) + [None] * (13 - len(row))
        product_name = _clean(values[1])
        if not product_name:
            continue

        data.append({
            "product_name": product_name,
            "article": _clean(values[2]),
            "sku": _clean(values[3]),
            "barcode": _clean(values[4]),
            "quantity": _number(values[7]),
            "sale_price": _number(values[8]),
            "revenue": _number(values[5]),
            "loyalty_payout": _number(values[6]),
            "return_amount": _number(values[9]),
            "return_quantity": _number(values[11]),
            "return_price": _number(values[12]),
            "return_loyalty": _number(values[10]),
        })

    return data


def _parse_ozon_orders_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as file:
        sample = file.read(8192)
        file.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(file, dialect=dialect)
        rows = []
        for row in reader:
            def get(name):
                return row.get(name, row.get("\ufeff" + name, ""))

            order_number = _clean(get("Номер заказа"))
            if not order_number:
                continue

            rows.append({
                "order_number": order_number,
                "shipment_number": _clean(get("Номер отправления")),
                "sale_date": _clean(get("Дата доставки")) or _clean(get("Дата отгрузки")),
                "status": _clean(get("Статус")),
                "product_name": _clean(get("Название товара")),
                "article": _clean(get("Артикул")),
                "sku": _clean(get("SKU")),
                "quantity": _number(get("Количество")),
                "sale_price": _number(get("Цена товара до скидок")),
                "revenue": _number(get("Сумма отправления")),
                "buyer_paid": _number(get("Оплачено покупателем")),
                "discount_percent": _number(get("Скидка %")),
                "discount_amount": _number(get("Скидка руб")),
            })
        return rows


def _import_rows(path, report_type, rows):
    file_hash = _file_hash(path)

    with get_connection() as connection:
        existing = connection.execute(
            "SELECT id FROM marketplace_imports WHERE file_hash = ?",
            (file_hash,)
        ).fetchone()
        if existing:
            return 0, 0, "Этот файл уже импортирован."

        cursor = connection.execute("""
            INSERT INTO marketplace_imports (
                marketplace, report_type, file_name, file_hash, rows_count
            ) VALUES (?, ?, ?, ?, 0)
        """, ("Ozon", report_type, path.name, file_hash))
        import_id = cursor.lastrowid
        added = 0
        duplicates = 0

        for index, row in enumerate(rows, start=1):
            article = row.get("article", "")
            product = _find_products(connection, article)
            product_id = product["id"] if product else None

            if report_type == "Реализация":
                external_key = (
                    f"{file_hash}:{index}:{article}:{row.get('product_name', '')}:"
                    f"{row.get('quantity', 0)}:{row.get('revenue', 0)}"
                )
                sale_date = None
                status = "Реализовано"
                order_number = None
                shipment_number = None
            else:
                external_key = (
                    f"{row.get('order_number')}:{row.get('shipment_number')}:"
                    f"{row.get('sku')}:{index}"
                )
                sale_date = row.get("sale_date")
                status = row.get("status")
                order_number = row.get("order_number")
                shipment_number = row.get("shipment_number")

            exists = connection.execute("""
                SELECT id FROM marketplace_sales
                WHERE marketplace = ? AND report_type = ? AND external_key = ?
            """, ("Ozon", report_type, external_key)).fetchone()

            if exists:
                duplicates += 1
                continue

            connection.execute("""
                INSERT INTO marketplace_sales (
                    marketplace, report_type, import_id, external_key,
                    order_number, shipment_number, sale_date, status,
                    product_name, article, sku, barcode, quantity,
                    sale_price, revenue, loyalty_payout, return_amount,
                    buyer_paid, discount_percent, discount_amount,
                    product_id, raw_data
                ) VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?
                )
            """, (
                "Ozon", report_type, import_id, external_key,
                order_number, shipment_number, sale_date, status,
                row.get("product_name", ""), article, row.get("sku", ""),
                row.get("barcode", ""), row.get("quantity", 0),
                row.get("sale_price", 0), row.get("revenue", 0),
                row.get("loyalty_payout", 0), row.get("return_amount", 0),
                row.get("buyer_paid", 0), row.get("discount_percent", 0),
                row.get("discount_amount", 0), product_id,
                repr(row),
            ))
            added += 1

        connection.execute(
            "UPDATE marketplace_imports SET rows_count = ? WHERE id = ?",
            (added, import_id)
        )

    return added, duplicates, None


class MarketplacePage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        init_marketplace_tables()
        self.setup_ui()
        self.load_data()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(35, 30, 35, 30)
        layout.setSpacing(14)

        title = QLabel("Маркетплейсы")
        title.setObjectName("pageTitle")
        layout.addWidget(title)

        subtitle = QLabel(
            "Импортируем данные Ozon из официальных выгрузок. "
            "На первом этапе сохраняются реализации и заказы отдельно, "
            "без выдумывания комиссии и логистики там, где их нет в файле."
        )
        subtitle.setObjectName("subtitle")
        subtitle.setWordWrap(True)
        layout.addWidget(subtitle)

        buttons = QHBoxLayout()
        realization_button = QPushButton("Импорт отчёта о реализации Ozon")
        realization_button.setObjectName("accentButton")
        realization_button.clicked.connect(self.import_realization)
        buttons.addWidget(realization_button)

        orders_button = QPushButton("Импорт заказов Ozon CSV")
        orders_button.clicked.connect(self.import_orders)
        buttons.addWidget(orders_button)

        refresh_button = QPushButton("Обновить")
        refresh_button.clicked.connect(self.load_data)
        buttons.addWidget(refresh_button)
        buttons.addStretch()
        layout.addLayout(buttons)

        self.summary = QLabel("")
        self.summary.setObjectName("pageTitle")
        layout.addWidget(self.summary)

        self.table = QTableWidget(0, 8)
        self.table.setHorizontalHeaderLabels([
            "Дата", "Тип", "Статус", "Артикул",
            "Товар", "Кол-во", "Выручка", "Товар в каталоге"
        ])
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(self.table, 1)

    def import_realization(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Выбери отчёт Ozon о реализации", str(Path.home()), "Excel (*.xlsx)"
        )
        if not path:
            return
        try:
            file_path = Path(path)
            rows = _parse_ozon_realization(file_path)
            added, duplicates, message = _import_rows(file_path, "Реализация", rows)
            if message:
                QMessageBox.information(self, "Импорт Ozon", message)
                return
            QMessageBox.information(
                self, "Импорт завершён",
                f"Добавлено строк: {added}\nПропущено дублей: {duplicates}\nВсего строк в файле: {len(rows)}"
            )
            self.load_data()
        except Exception as error:
            QMessageBox.critical(self, "Ошибка импорта Ozon", f"Не удалось импортировать отчёт:\n{error}")

    def import_orders(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Выбери выгрузку заказов Ozon", str(Path.home()), "CSV (*.csv)"
        )
        if not path:
            return
        try:
            file_path = Path(path)
            rows = _parse_ozon_orders_csv(file_path)
            added, duplicates, message = _import_rows(file_path, "Заказы", rows)
            if message:
                QMessageBox.information(self, "Импорт Ozon", message)
                return
            QMessageBox.information(
                self, "Импорт завершён",
                f"Добавлено строк: {added}\nПропущено дублей: {duplicates}\nВсего строк в файле: {len(rows)}"
            )
            self.load_data()
        except Exception as error:
            QMessageBox.critical(self, "Ошибка импорта Ozon", f"Не удалось импортировать CSV:\n{error}")

    def load_data(self):
        with get_connection() as connection:
            rows = connection.execute("""
                SELECT sale_date, report_type, status, article,
                       product_name, quantity, revenue, product_id
                FROM marketplace_sales
                ORDER BY id DESC
                LIMIT 500
            """).fetchall()
            realization = connection.execute("""
                SELECT COALESCE(SUM(revenue), 0) AS revenue,
                       COALESCE(SUM(quantity), 0) AS quantity,
                       COUNT(*) AS rows_count
                FROM marketplace_sales
                WHERE marketplace = 'Ozon' AND report_type = 'Реализация'
            """).fetchone()
            orders = connection.execute("""
                SELECT COUNT(*) AS rows_count
                FROM marketplace_sales
                WHERE marketplace = 'Ozon' AND report_type = 'Заказы'
            """).fetchone()

        self.summary.setText(
            f"Ozon: реализация {money(realization['revenue'])} | "
            f"товаров {int(realization['quantity'] or 0)} | "
            f"строк реализации {realization['rows_count']} | "
            f"строк заказов {orders['rows_count']}"
        )
        self.table.setRowCount(len(rows))
        for index, row in enumerate(rows):
            values = (
                row["sale_date"] or "—", row["report_type"] or "",
                row["status"] or "", row["article"] or "—",
                row["product_name"] or "", str(row["quantity"] or 0),
                money(row["revenue"]), "Да" if row["product_id"] else "Нет",
            )
            for column, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                if column in (5, 6):
                    item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(index, column, item)
